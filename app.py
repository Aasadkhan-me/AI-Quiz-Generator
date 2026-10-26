
"""Notes & Quiz Generator — AI-powered study notes and interactive quizzes."""

import io
import json
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pydantic import BaseModel
from pypdf import PdfReader


# ---------- Configuration ----------

DEFAULT_MODEL = "gemini-3.8-flash"
CHUNK_CHARS = 30_000
MAX_CHUNKS = 25
QUIZ_DIRECT_LIMIT = 150_000


# ---------- Quiz schema ----------

class Question(BaseModel):
    question: str
    options: list[str]
    answer_index: int
    explanation: str


# ---------- File handling ----------

def extract_text(uploaded) -> str:
    """Extract text from PDF, DOCX, TXT, or Markdown files."""

    name = uploaded.name.lower()
    data = uploaded.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(
            page.extract_text() or "" for page in reader.pages
        )

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]

        # Include text from tables too.
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(paragraphs)

    return data.decode("utf-8", errors="ignore")


def chunk_text(text: str, size: int = CHUNK_CHARS) -> list[str]:
    """Split long text into manageable chunks."""

    chunks = []
    current = ""

    for paragraph in text.splitlines():
        if len(current) + len(paragraph) + 1 > size and current:
            chunks.append(current)
            current = ""

        while len(paragraph) > size:
            chunks.append(paragraph[:size])
            paragraph = paragraph[size:]

        current += paragraph + "\n"

    if current.strip():
        chunks.append(current)

    return chunks


# ---------- Gemini API with retries ----------

def call_gemini(
    client,
    model,
    prompt,
    system,
    schema=None,
    retries=5,
):
    """Call Gemini and retry temporary service failures."""

    config = {
        "system_instruction": system,
        "temperature": 0.4,
    }

    if schema is not None:
        config["response_mime_type"] = "application/json"
        config["response_schema"] = schema

    for attempt in range(retries):
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt,
                config=types.GenerateContentConfig(**config),
            )

            if not response.text:
                raise ValueError(
                    "Gemini returned an empty response. "
                    "Try again or use a smaller document."
                )

            return response.text

        except Exception as error:
            message = str(error)
            upper = message.upper()

            # Permanent errors should not be retried.
            if (
                "NOT_FOUND" in upper
                or "404" in upper
                or "INVALID_ARGUMENT" in upper
                or "400" in upper
                or "UNAUTHENTICATED" in upper
                or "401" in upper
                or "PERMISSION_DENIED" in upper
                or "403" in upper
            ):
                raise RuntimeError(
                    f"Gemini configuration or access error: {message}"
                ) from error

            retryable = any(
                term in upper
                for term in (
                    "503",
                    "UNAVAILABLE",
                    "429",
                    "RESOURCE_EXHAUSTED",
                    "500",
                    "INTERNAL",
                    "504",
                    "DEADLINE_EXCEEDED",
                    "TIMEOUT",
                    "CONNECTION",
                )
            )

            if not retryable or attempt == retries - 1:
                raise RuntimeError(
                    f"Gemini request failed after "
                    f"{attempt + 1} attempt(s): {message}"
                ) from error

            wait_seconds = min(2 ** (attempt + 1), 30)

            st.warning(
                f"Gemini is temporarily busy. Retrying in "
                f"{wait_seconds} seconds "
                f"({attempt + 2}/{retries})..."
            )

            time.sleep(wait_seconds)

    raise RuntimeError("Gemini could not complete the request.")


# ---------- Notes generation ----------

NOTES_SYSTEM = """
You are an expert tutor and educational content writer.

Create accurate, clear, well-organized study notes using only
the supplied educational material.

Do not invent facts or claim unsupported information.
Explain difficult concepts in simple language.
Preserve important definitions, formulas, examples, and facts.
"""


def generate_notes(client, model, chunks, style, progress=None):
    """Generate study notes for each document chunk."""

    parts = []

    for index, chunk in enumerate(chunks, start=1):
        prompt = f"""
Create {style} study notes in Markdown.

Requirements:
- Use clear headings and subheadings.
- Explain important concepts in simple language.
- Use bullet points and bold key terms.
- Preserve important definitions, formulas, and examples.
- Include a Key Takeaways section.
- Include a Key Terms glossary when relevant.
- Do not invent information absent from the source.

Document part {index} of {len(chunks)}:

{chunk}
"""

        result = call_gemini(
            client=client,
            model=model,
            prompt=prompt,
            system=NOTES_SYSTEM,
        )

        parts.append(result)

        if progress:
            progress(index / len(chunks))

    return "\n\n---\n\n".join(parts)


# ---------- Quiz generation ----------

def validate_questions(raw: str) -> list[dict]:
    """Validate quiz questions returned by Gemini."""

    try:
        items = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return []

    if not isinstance(items, list):
        return []

    valid_questions = []

    for item in items:
        try:
            question = Question(**item)
        except Exception:
            continue

        # Require four unique options and a valid answer index.
        if len(question.options) != 4:
            continue

        if len(set(question.options)) != 4:
            continue

        if not 0 <= question.answer_index < 4:
            continue

        valid_questions.append(question.model_dump())

    return valid_questions


def generate_quiz(client, model, source, number, difficulty):
    """Generate multiple-choice questions."""

    prompt = f"""
Create exactly {number} multiple-choice questions at
{difficulty} difficulty based only on the provided material.

Requirements:
- Exactly four options per question.
- Exactly one correct answer.
- answer_index must be zero-based: 0, 1, 2, or 3.
- Include a short explanation for every answer.
- Cover different sections and important concepts.
- Avoid duplicate or trivial questions.
- Do not use "all of the above".
- Return the questions in the required JSON structure.

MATERIAL:

{source}
"""

    raw = call_gemini(
        client=client,
        model=model,
        prompt=prompt,
        system=NOTES_SYSTEM,
        schema=list[Question],
    )

    return validate_questions(raw)


# ---------- Streamlit application ----------

def main():
    st.set_page_config(
        page_title="Notes & Quiz Generator",
        page_icon="📚",
        layout="wide",
    )

    st.title("📚 Notes & Quiz Generator")
    st.write(
        "Upload a chapter or book and generate professional "
        "study notes and an interactive practice quiz."
    )

    # Settings
    with st.sidebar:
        st.header("⚙️ Settings")

        try:
            saved_key = st.secrets.get("GEMINI_API_KEY", "")
        except Exception:
            saved_key = ""

        api_key = st.text_input(
            "Gemini API key",
            value=saved_key,
            type="password",
            help="Get an API key from Google AI Studio.",
        )

        model = st.text_input(
            "Gemini model",
            value=DEFAULT_MODEL,
            help="Enter a model ID available to your API key.",
        ).strip()

        style = st.selectbox(
            "Notes style",
            ["concise", "detailed", "exam-focused"],
            index=1,
        )

        number = st.slider(
            "Number of quiz questions",
            min_value=5,
            max_value=30,
            value=10,
        )

        difficulty = st.selectbox(
            "Quiz difficulty",
            ["easy", "medium", "hard"],
            index=1,
        )

    # Upload or paste source text
    uploaded = st.file_uploader(
        "📄 Upload your study material",
        type=["pdf", "docx", "txt", "md"],
    )

    pasted = st.text_area(
        "Or paste your chapter text here",
        height=150,
        placeholder="Paste at least 200 characters...",
    )

    # Generate
    if st.button("🚀 Generate Notes & Quiz", type="primary"):
        if not api_key.strip():
            st.error("Please enter your Gemini API key in the sidebar.")
            st.stop()

        if not model:
            st.error("Please enter a Gemini model name.")
            st.stop()

        if not uploaded and not pasted.strip():
            st.error("Please upload a file or paste some text.")
            st.stop()

        try:
            source_text = (
                extract_text(uploaded)
                if uploaded
                else pasted
            )
        except Exception as error:
            st.error(f"Could not read the uploaded file: {error}")
            st.stop()

        if len(source_text.strip()) < 200:
            st.error(
                "Not enough text was extracted. "
                "Scanned PDFs without selectable text are not supported."
            )
            st.stop()

        chunks = chunk_text(source_text)

        if len(chunks) > MAX_CHUNKS:
            st.warning(
                f"The document is large. Only the first "
                f"{MAX_CHUNKS} chunks will be processed."
            )
            chunks = chunks[:MAX_CHUNKS]

        client = genai.Client(api_key=api_key.strip())

        progress = st.progress(
            0.0,
            text="Preparing your study notes...",
        )

        try:
            notes = generate_notes(
                client=client,
                model=model,
                chunks=chunks,
                style=style,
                progress=lambda value: progress.progress(
                    value,
                    text="Generating study notes...",
                ),
            )

            progress.progress(
                1.0,
                text="Creating your practice quiz...",
            )

            quiz_source = "\n\n".join(chunks)

            # Use generated notes for very large quiz sources.
            if len(quiz_source) > QUIZ_DIRECT_LIMIT:
                quiz_source = notes

            quiz = generate_quiz(
                client=client,
                model=model,
                source=quiz_source,
                number=number,
                difficulty=difficulty,
            )

            # Save results only after successful generation.
            st.session_state["notes"] = notes
            st.session_state["quiz"] = quiz
            st.session_state["qid"] = (
                st.session_state.get("qid", 0) + 1
            )
            st.session_state["submitted"] = False
            st.session_state.pop("answers", None)

            st.success("Your notes and quiz are ready!")

        except Exception as error:
            st.error(str(error))
            st.info(
                "If Gemini reports 503 or 429, wait and try again. "
                "Repeated requests may also encounter API quota limits."
            )
            st.stop()

        finally:
            progress.empty()

    # Show results
    if "notes" not in st.session_state:
        st.info(
            "Your generated notes and quiz will appear here."
        )
        return

    notes_tab, quiz_tab = st.tabs(
        ["📝 Study Notes", "❓ Interactive Quiz"]
    )

    # Notes tab
    with notes_tab:
        st.markdown(st.session_state["notes"])

        st.download_button(
            "⬇️ Download Notes (.md)",
            data=st.session_state["notes"],
            file_name="study_notes.md",
            mime="text/markdown",
        )

    # Quiz tab
    with quiz_tab:
        quiz = st.session_state["quiz"]

        if not quiz:
            st.warning(
                "No valid questions were generated. "
                "Please try generating the quiz again."
            )
            return

        quiz_id = st.session_state["qid"]

        with st.form(f"quiz_form_{quiz_id}"):
            answers = []

            for index, question in enumerate(quiz):
                selected = st.radio(
                    f"Q{index + 1}. {question['question']}",
                    options=question["options"],
                    index=None,
                    key=f"answer_{quiz_id}_{index}",
                )

                answers.append(selected)

            submitted = st.form_submit_button("Check My Answers")

        if submitted:
            st.session_state["submitted"] = True
            st.session_state["answers"] = answers

        if st.session_state.get("submitted"):
            saved_answers = st.session_state["answers"]
            score = 0

            st.subheader("📊 Quiz Results")

            for index, question in enumerate(quiz):
                correct = question["options"][
                    question["answer_index"]
                ]
                selected = saved_answers[index]

                if selected == correct:
                    score += 1
                    st.success(f"Q{index + 1}: Correct!")
                else:
                    st.error(f"Q{index + 1}: Incorrect.")

                if selected is None:
                    st.write("Your answer: Not answered")
                else:
                    st.write(f"Your answer: {selected}")

                st.write(f"Correct answer: {correct}")
                st.write(
                    f"**Explanation:** {question['explanation']}"
                )

            st.metric(
                "Final Score",
                f"{score} / {len(quiz)}",
            )

            st.progress(score / len(quiz))

        st.download_button(
            "⬇️ Download Quiz (.json)",
            data=json.dumps(
                quiz,
                indent=2,
                ensure_ascii=False,
            ),
            file_name="practice_quiz.json",
            mime="application/json",
        )


if __name__ == "__main__":
    main()


